from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .pdf import _donnees_rapport_comptable

ENTETE_FONT = Font(bold=True, color='FFFFFF')
ENTETE_FILL = PatternFill('solid', fgColor='1A1440')
FORMAT_MONETAIRE = '#,##0.00'


def _ecrire_entete(feuille, colonnes):
    feuille.append(colonnes)
    for cellule in feuille[1]:
        cellule.font = ENTETE_FONT
        cellule.fill = ENTETE_FILL
        cellule.alignment = Alignment(horizontal='center')


def _ajuster_largeurs(feuille):
    for colonne in feuille.columns:
        longueur = max((len(str(c.value)) for c in colonne if c.value is not None), default=10)
        feuille.column_dimensions[get_column_letter(colonne[0].column)].width = min(longueur + 2, 40)


def _feuille_ventes(feuille, ventes):
    _ecrire_entete(feuille, ['Date', 'Numéro', 'Client', 'Sous-total', 'TPS', 'TVQ', 'Total'])
    for vente in ventes:
        feuille.append([
            vente.date_creation.date(), vente.numero, vente.client.nom_entreprise,
            float(vente.sous_total), float(vente.montant_tps), float(vente.montant_tvq), float(vente.total),
        ])
    for ligne in feuille.iter_rows(min_row=2, min_col=4, max_col=7):
        for cellule in ligne:
            cellule.number_format = FORMAT_MONETAIRE
    _ajuster_largeurs(feuille)


def _feuille_depenses(feuille, depenses):
    _ecrire_entete(
        feuille, ['Date', 'Fournisseur', 'Description', 'Compte de grand livre', 'Sous-total', 'TPS', 'TVQ', 'Total'],
    )
    for depense in depenses:
        feuille.append([
            depense.date, depense.fournisseur, depense.description, depense.compte_grand_livre.nom,
            float(depense.sous_total), float(depense.tps), float(depense.tvq), float(depense.total),
        ])
    for ligne in feuille.iter_rows(min_row=2, min_col=5, max_col=8):
        for cellule in ligne:
            cellule.number_format = FORMAT_MONETAIRE
    _ajuster_largeurs(feuille)


def _feuille_sommaire(feuille, annee, trimestre, d):
    feuille.append([f'Rapport comptable — T{trimestre} {annee}'])
    feuille['A1'].font = Font(bold=True, size=14)
    feuille.append([])
    _ecrire_entete(feuille, ['', 'Montant'])
    lignes = [
        ('Sous-total des ventes', d['sous_total_ventes']),
        ('TPS perçue', d['tps_percue']),
        ('TVQ perçue', d['tvq_percue']),
        ('Total des ventes', d['total_ventes']),
        ('', None),
        ('Sous-total des dépenses', d['sous_total_depenses']),
        ('TPS payée', d['tps_payee']),
        ('TVQ payée', d['tvq_payee']),
        ('Total des dépenses', d['total_depenses']),
        ('', None),
        ('Solde TPS (perçue − payée)', d['tps_percue'] - d['tps_payee']),
        ('Solde TVQ (perçue − payée)', d['tvq_percue'] - d['tvq_payee']),
    ]
    for libelle, montant in lignes:
        feuille.append([libelle, float(montant) if montant is not None else None])
    for ligne in feuille.iter_rows(min_row=4, min_col=2, max_col=2):
        for cellule in ligne:
            if cellule.value is not None:
                cellule.number_format = FORMAT_MONETAIRE
    _ajuster_largeurs(feuille)


def generer_excel_rapport_comptable(annee, trimestre):
    d = _donnees_rapport_comptable(annee, trimestre)

    classeur = Workbook()
    feuille_ventes = classeur.active
    feuille_ventes.title = 'Ventes'
    _feuille_ventes(feuille_ventes, d['ventes'])

    _feuille_depenses(classeur.create_sheet('Dépenses'), d['depenses'])
    _feuille_sommaire(classeur.create_sheet('Sommaire'), annee, trimestre, d)

    tampon = BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()
